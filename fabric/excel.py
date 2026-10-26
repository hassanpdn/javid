"""The three Excel files a traditional wholesaler keeps, and everything the app does with them.

build()       makes a believable four months of business in three workbooks (stock, sales, accounts),
              with a few situations planted for the demo (see PLANTED), then reads them in.
import_all()  reads the three workbooks into fabric.db. Like a real office's files they are a little messy
              (a customer written three ways, Persian digits, a hand-typed total row); every fix is noted
              in import_notes so the panel can show what the assistant cleaned up.
              PLANTED also leaves the marks a theft leaves: an invoice missing from the sales file and a day
              when less cash reached the bank. facts.reconcile() finds them from the numbers alone.
add_invoice() a sale confirmed at the phone desk goes into all three files: a row in sales, fewer rolls in
              stock, and the cheque in accounts when it is paid by cheque.
"""
import os, random, re
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from jalali import jstr, to_jalali, to_gregorian
from . import store, business

DIR = os.getenv("FABRIC_DIR", "fabric_files")
FILES = {   # key: (file on disk, name shown and used for downloads)
    "stock": ("stock.xlsx", "انبار.xlsx"),
    "sales": ("sales.xlsx", "فروش.xlsx"),
    "accounts": ("accounts.xlsx", "حسابداری.xlsx"),
}

def today():
    return date.fromisoformat(os.getenv("DEMO_TODAY") or date.today().isoformat())

def path(key):
    return os.path.join(DIR, FILES[key][0])

# ---------- column names, the same ones a person would type ----------

STOCK_COLS = ["کد کالا", "شرح کالا", "جنس", "رنگ", "عرض (سانتی‌متر)", "تعداد توپ", "متراژ هر توپ", "متراژ موجود",
              "موجودی اول دوره (متر)", "قیمت خرید هر متر", "قیمت فروش هر متر", "تأمین‌کننده", "تاریخ آخرین ورود"]
STOCK_IN_COLS = ["تاریخ", "کد کالا", "شرح کالا", "توپ", "متراژ", "تأمین‌کننده"]
SALES_COLS = ["شماره فاکتور", "تاریخ", "مشتری", "کد کالا", "شرح کالا", "توپ", "متراژ", "قیمت هر متر", "مبلغ",
              "نوع پرداخت", "فروشنده"]
CUSTOMER_COLS = ["کد مشتری", "نام", "فروشگاه", "شهر", "تلفن", "سقف اعتبار", "مانده اول دوره"]
RECEIPT_COLS = ["تاریخ", "مشتری", "مبلغ", "روش"]
CHEQUE_IN_COLS = ["شماره چک", "بانک", "مشتری", "مبلغ", "تاریخ سررسید", "وضعیت"]
CHEQUE_OUT_COLS = ["شماره چک", "در وجه", "مبلغ", "تاریخ سررسید", "وضعیت"]
BANK_COLS = ["تاریخ", "موجودی حساب"]
WEB_COLS = ["شماره سفارش", "تاریخ", "ساعت", "مشتری", "تلفن", "کد کالا", "شرح کالا", "توپ", "متراژ", "قیمت هر متر", "مبلغ", "وضعیت پرداخت"]
WEB_SHEET = "سفارش‌های سایت"
CASH_COLS = ["تاریخ", "واریز فروش نقدی"]
MONEY_COLS = {"قیمت خرید هر متر", "قیمت فروش هر متر", "قیمت هر متر", "مبلغ", "سقف اعتبار", "مانده اول دوره", "موجودی حساب", "واریز فروش نقدی"}

# ---------- the simulated business ----------

# (name, material, colors, width, meters per roll, sell price per meter in toman, supplier)
CATALOG = [
    ("کتان کش", "کتان", ["آبی نفتی", "مشکی", "کرم"], 150, 50, 185_000, "نساجی کاشان"),
    ("کتان ساده", "کتان", ["سفید", "آبی آسمانی"], 150, 50, 160_000, "نساجی کاشان"),
    ("مخمل کبریتی", "مخمل", ["زرشکی", "قهوه‌ای"], 145, 40, 240_000, "بازرگانی آسیا"),
    ("مخمل ساده", "مخمل", ["سرمه‌ای", "سبز یشمی"], 145, 40, 260_000, "بازرگانی آسیا"),
    ("حریر", "حریر", ["سفید", "صورتی"], 150, 60, 120_000, "بازرگانی آسیا"),
    ("کرپ مازراتی", "کرپ", ["مشکی", "کرم", "طوسی"], 150, 50, 210_000, "پارچه یزدباف"),
    ("لینن", "لینن", ["بژ", "سفید"], 145, 45, 290_000, "بازرگانی آسیا"),
    ("ساتن آمریکایی", "ساتن", ["مشکی", "زرشکی"], 150, 50, 150_000, "پارچه یزدباف"),
    ("چادری کن‌کن", "چادری", ["مشکی"], 115, 60, 420_000, "بازرگانی آسیا"),
    ("چادری حریر اسود", "چادری", ["مشکی"], 115, 60, 380_000, "بازرگانی آسیا"),
    ("فاستونی", "فاستونی", ["طوسی", "سرمه‌ای"], 150, 40, 520_000, "ریسندگی بروجرد"),
    ("ویسکوز", "ویسکوز", ["گلدار", "سبز"], 145, 50, 175_000, "پارچه یزدباف"),
    ("جین", "جین", ["آبی", "مشکی"], 150, 45, 230_000, "نساجی کاشان"),
    ("مخمل پرده‌ای", "پرده‌ای", ["طلایی", "طوسی"], 280, 40, 340_000, "بازرگانی آسیا"),
    ("ترگال", "ترگال", ["سرمه‌ای", "مشکی"], 150, 50, 140_000, "ریسندگی بروجرد"),
    ("پشمی پالتویی", "پشمی", ["شتری", "ذغالی"], 150, 35, 610_000, "ریسندگی بروجرد"),
    ("ململ", "ململ", ["سفید"], 150, 60, 95_000, "پارچه یزدباف"),
]
# (name, shop, city)
CUSTOMERS = [
    ("محمود رضایی", "پارچه‌سرای رضایی", "مشهد"), ("علی کریمی", "تولیدی پوشاک کریمی", "تبریز"),
    ("مهدی فرهادی", "مزون فرهادی", "شیراز"), ("رضا احمدی", "پارچه احمدی", "اصفهان"),
    ("حسین موسوی", "خیاطی موسوی", "قم"), ("مریم صالحی", "مزون صالحی", "تهران"),
    ("جواد قاسمی", "پارچه‌فروشی قاسمی", "اهواز"), ("سعید جعفری", "تولیدی جعفری", "کرج"),
    ("امیر حسینی", "پارچه حسینی", "رشت"), ("ناصر عباسی", "پارچه عباسی", "کرمانشاه"),
    ("کاظم یزدانی", "پارچه یزدانی", "یزد"), ("فاطمه نوری", "مزون نوری", "اصفهان"),
    ("غلامرضا شریفی", "پارچه شریفی", "زاهدان"), ("بهروز ملکی", "پارچه ملکی", "ارومیه"),
    ("یوسف بلوچی", "پارچه بلوچ", "بندرعباس"), ("محسن طاهری", "تولیدی طاهری", "مشهد"),
    ("مجید رحیمی", "پارچه رحیمی", "همدان"), ("داوود کاظم‌زاده", "پارچه کاظم‌زاده", "کرمان"),
    ("عباس نجفی", "پارچه نجفی", "اراک"), ("حمیدرضا صفری", "تولیدی صفری", "قزوین"),
    ("اکبر میرزایی", "پارچه میرزایی", "سنندج"), ("مرتضی اکبری", "پارچه اکبری", "ساری"),
    ("علیرضا ابراهیمی", "پارچه ابراهیمی", "بوشهر"), ("رسول زارعی", "تولیدی زارعی", "شیراز"),
]
BANKS = ["ملت", "ملی", "صادرات", "تجارت", "سپه", "پاسارگاد", "رسالت"]
# Situations the demo shows. They are found again by facts.py from the data alone, not from these names.
PLANTED = {
    "bounced": "محمود رضایی",      # big regular buyer with a bounced cheque, still ordering
    "over_limit": "علی کریمی",     # owes more than his credit limit allows
    "quiet": "مهدی فرهادی",        # bought every week or two, nothing for two months
    "fast": ("کتان کش", "آبی نفتی"),  # sells fast, about three days of stock left
    "dead": [("پشمی پالتویی", "شتری"), ("مخمل پرده‌ای", "طلایی"), ("ساتن آمریکایی", "زرشکی"), ("ململ", "سفید")],
    # what the three files show after a theft: an invoice that was written and then removed from the sales file
    # (its fabric left the warehouse, so stock no longer adds up), and one day when less cash reached the bank
    "hidden_invoice": {"days_ago": 12, "customer": "رضا احمدی",
                       "lines": [(("مخمل ساده", "سرمه‌ای"), 2), (("چادری کن‌کن", "مشکی"), 1)]},
    "cash_short": {"days_ago": 5, "amount": 6_500_000},
}

def _round(v, step=1000):
    return int(round(v / step) * step)

def simulate():
    """The business's last four months, as rows for the three workbooks."""
    t = today()
    rnd = random.Random(t.toordinal())
    seller = business.current()["seller_name"]

    products = []
    for i, (name, material, colors, width, roll_m, price, supplier) in enumerate(CATALOG):
        for j, color in enumerate(colors):
            products.append({"code": f"P-{101 + len(products)}", "name": name, "material": material, "color": color,
                             "width": width, "roll_meters": roll_m, "sell": price, "buy": _round(price * .78),
                             "supplier": supplier, "weight": rnd.choice([1, 2, 3, 4])})
    by_key = {(p["name"], p["color"]): p for p in products}
    fast = by_key[PLANTED["fast"]]
    fast["weight"] = 8
    dead = [by_key[k] for k in PLANTED["dead"]]
    for p in dead:
        p["weight"] = 0

    customers = [{"id": i + 1, "name": n, "shop": s, "city": c, "phone": f"09{rnd.randint(11, 39)}{rnd.randint(1000000, 9999999)}",
                  "weight": rnd.choice([1, 1, 2, 2, 3])} for i, (n, s, c) in enumerate(CUSTOMERS)]
    cust = {c["name"]: c for c in customers}
    cust[PLANTED["bounced"]]["weight"] = 6
    cust[PLANTED["over_limit"]]["weight"] = 4
    cust[PLANTED["quiet"]]["weight"] = 2

    start = t - timedelta(days=120)
    sales, cheques_in, receipts = [], [], []
    number = 1400
    day = start
    hidden = PLANTED["hidden_invoice"]
    hidden_done = False
    while day <= t:
        wd = day.weekday()                       # Monday = 0 ... Friday = 4 (closed)
        if wd not in (3, 4) and not hidden_done and (t - day).days <= hidden["days_ago"]:
            # the invoice that will be missing from the sales file: written first thing that day, paid in cash
            hidden_done = True
            number += 1
            for key, rolls in hidden["lines"]:
                p = by_key[key]
                sales.append({"no": number, "date": day, "customer": hidden["customer"], "p": p, "rolls": rolls,
                              "meters": rolls * p["roll_meters"], "price": p["sell"], "amount": _round(rolls * p["roll_meters"] * p["sell"]),
                              "payment": "نقد", "seller": seller, "hidden": True})
        if wd != 4:
            count = rnd.randint(1, 3) if wd == 3 else rnd.randint(3, 6)
            if day == t:
                count = 2                        # this morning so far
            for _ in range(count):
                pool = [c for c in customers if not (c["name"] == PLANTED["quiet"] and (t - day).days < 65)]
                c = rnd.choices(pool, [x["weight"] for x in pool])[0]
                big = c["name"] in (PLANTED["bounced"], PLANTED["over_limit"])
                live = [p for p in products if p["weight"]]
                lines = []
                for p in rnd.sample(rnd.choices(live, [p["weight"] for p in live], k=6), k=rnd.randint(1, 3)):
                    if any(l["p"] is p for l in lines):
                        continue
                    if rnd.random() < .15 and not big:  # a cut from a roll instead of whole rolls
                        lines.append({"p": p, "rolls": 0, "meters": rnd.choice([10, 15, 20, 25, 30])})
                    else:
                        rolls = rnd.randint(2, 6) if big else rnd.randint(1, 4)
                        lines.append({"p": p, "rolls": rolls, "meters": rolls * p["roll_meters"]})
                total = sum(_round(l["meters"] * l["p"]["sell"]) for l in lines)
                if c["name"] == PLANTED["over_limit"]:
                    payment = rnd.choices(["نسیه", "چک"], [.6, .4])[0]
                else:
                    payment = rnd.choices(["نقد", "چک", "نسیه"], [.25, .5, .25])[0]
                number += 1
                for l in lines:
                    sales.append({"no": number, "date": day, "customer": c["name"], "p": l["p"], "rolls": l["rolls"],
                                  "meters": l["meters"], "price": l["p"]["sell"], "amount": _round(l["meters"] * l["p"]["sell"]),
                                  "payment": payment, "seller": seller})
                if payment == "چک":
                    due = day + timedelta(days=rnd.choice([30, 45, 60, 90]))
                    cheques_in.append({"number": str(rnd.randint(100000, 999999)), "bank": rnd.choice(BANKS),
                                       "customer": c["name"], "amount": total, "due": due,
                                       "status": "وصول شد" if due < t else "در جریان"})
        day += timedelta(days=1)

    # the bounced cheque: the big Mashhad buyer's cheque came back nine days ago and is still unpaid
    cheques_in.append({"number": "748215", "bank": "ملت", "customer": PLANTED["bounced"], "amount": 128_000_000,
                       "due": t - timedelta(days=9), "status": "برگشتی"})

    # customers pay their credit purchases down from time to time; the over-limit customer pays little
    for c in customers:
        credit = [s for s in sales if s["customer"] == c["name"] and s["payment"] == "نسیه"]
        owed = sum(s["amount"] for s in credit)
        share = .25 if c["name"] == PLANTED["over_limit"] else rnd.uniform(.65, .9)
        left = _round(owed * share, 1_000_000)
        dates = sorted({s["date"] for s in credit})
        while left > 0 and dates:
            d = min(dates[rnd.randrange(len(dates))] + timedelta(days=rnd.randint(7, 25)), t - timedelta(days=1))
            amount = min(left, _round(rnd.uniform(.2, .5) * owed, 1_000_000) or 1_000_000)
            receipts.append({"date": d, "customer": c["name"], "amount": amount, "method": rnd.choice(["واریز", "کارت به کارت", "نقد"])})
            left -= amount
        c["opening"] = rnd.choice([0, 0, 0, _round(rnd.uniform(5, 40) * 1e6, 1_000_000)])
    receipts.sort(key=lambda r: r["date"])

    # credit limits: room to spare for everyone except the over-limit customer
    for c in customers:
        debt = c["opening"] + sum(s["amount"] for s in sales if s["customer"] == c["name"] and s["payment"] == "نسیه") \
            - sum(r["amount"] for r in receipts if r["customer"] == c["name"])
        debt += sum(q["amount"] for q in cheques_in if q["customer"] == c["name"] and q["status"] == "برگشتی")
        open_cheques = sum(q["amount"] for q in cheques_in if q["customer"] == c["name"] and q["status"] == "در جریان")
        exposure = max(debt, 0) + open_cheques
        if c["name"] == PLANTED["over_limit"]:
            c["limit"] = max(50_000_000, _round(exposure * .78, 50_000_000))
        else:
            c["limit"] = max(150_000_000, _round(exposure * 1.6 + 50_000_000, 50_000_000))

    # stock today: what is left on the shelves. The fast seller is almost gone; the dead stock sits there.
    sold_30 = {}
    for s in sales:
        if (t - s["date"]).days < 30:
            sold_30[s["p"]["code"]] = sold_30.get(s["p"]["code"], 0) + s["meters"]
    for p in products:   # enough for two to five weeks at the current pace
        per_day = sold_30.get(p["code"], 0) / 26
        p["rolls"] = max(rnd.randint(4, 12), -(-int(per_day * rnd.randint(14, 35)) // p["roll_meters"]))
        p["received"] = t - timedelta(days=rnd.randint(10, 60))
    fast["rolls"] = max(1, round(sold_30.get(fast["code"], 0) / 26 * 3 / fast["roll_meters"]))   # about three days left
    for p in dead:
        p["rolls"] = rnd.randint(12, 20)
        p["received"] = t - timedelta(days=rnd.randint(150, 220))
    for p in products:
        p["meters"] = p["rolls"] * p["roll_meters"]

    # deliveries from suppliers over the four months, and the stock at the start of the period, chosen so that
    # start + deliveries - everything that really left = what is on the shelves now
    stock_in = []
    for p in products:
        left = sum(x["meters"] for x in sales if x["p"] is p)
        rolls_in = min(int(left * rnd.uniform(.65, .95) / p["roll_meters"]), int((left + p["meters"]) / p["roll_meters"]))
        p["opening"] = p["meters"] + left - rolls_in * p["roll_meters"]
        parts = sorted(rnd.sample(range(8, 110), k=min(3, rolls_in)), reverse=True) if rolls_in else []
        for i, ago in enumerate(parts):
            n = rolls_in // len(parts) + (1 if i < rolls_in % len(parts) else 0)
            stock_in.append({"date": t - timedelta(days=ago), "p": p, "rolls": n, "meters": n * p["roll_meters"]})
        if parts:
            p["received"] = t - timedelta(days=parts[-1])
    stock_in.sort(key=lambda r: r["date"])

    # cash sales are taken to the bank every evening; one evening less arrived than was sold
    cash = {}
    for x in sales:
        if x["payment"] == "نقد" and not x.get("hidden") and x["date"] < t:
            cash[x["date"]] = cash.get(x["date"], 0) + x["amount"]
    short = PLANTED["cash_short"]
    short_day = min((d for d in cash if cash[d] > short["amount"] * 2), key=lambda d: abs((t - d).days - short["days_ago"]))
    deposits = [{"date": d, "amount": v - (short["amount"] if d == short_day else 0)} for d, v in sorted(cash.items())]

    # the wholesaler's own cheques to suppliers. Next week more goes out than comes in plus what is in the bank.
    incoming_week = sum(q["amount"] for q in cheques_in if q["status"] == "در جریان" and 0 <= (q["due"] - t).days <= 7)
    bank = 210_000_000
    cheques_out = []
    for k in range(1, 5):   # one per month, already paid
        cheques_out.append({"number": str(rnd.randint(100000, 999999)), "payee": rnd.choice(["نساجی کاشان", "بازرگانی آسیا", "پارچه یزدباف"]),
                            "amount": _round(rnd.uniform(250, 450) * 1e6, 1_000_000), "due": t - timedelta(days=30 * k - 5), "status": "پاس شد"})
    gap = _round((bank + incoming_week) * 1.18, 1_000_000)
    cheques_out.append({"number": str(rnd.randint(100000, 999999)), "payee": "بازرگانی آسیا", "amount": _round(gap * .6, 1_000_000),
                        "due": t + timedelta(days=3), "status": "در جریان"})
    cheques_out.append({"number": str(rnd.randint(100000, 999999)), "payee": "نساجی کاشان", "amount": _round(gap * .4, 1_000_000),
                        "due": t + timedelta(days=6), "status": "در جریان"})
    cheques_out.append({"number": str(rnd.randint(100000, 999999)), "payee": "ریسندگی بروجرد", "amount": 180_000_000,
                        "due": t + timedelta(days=24), "status": "در جریان"})
    return {"products": products, "customers": customers, "sales": sales, "receipts": receipts,
            "cheques_in": cheques_in, "cheques_out": cheques_out, "bank": bank, "stock_in": stock_in, "deposits": deposits}

# ---------- writing the workbooks ----------

HEAD_FILL = PatternFill("solid", fgColor="1F6F43")

def _sheet(ws, cols, rows):
    """A sheet the way an office keeps it: right to left, a coloured header row, numbers with separators."""
    ws.sheet_view.rightToLeft = True
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    for i, col in enumerate(cols, 1):
        letter = get_column_letter(i)
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200]]) + 4
        ws.column_dimensions[letter].width = min(max(width, 10), 34)
        if col in MONEY_COLS:
            for cell in ws[letter][1:]:
                cell.number_format = "#,##0"
    ws.freeze_panes = "A2"

def _fa(text):
    return str(text).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))

def write_workbooks(data):
    os.makedirs(DIR, exist_ok=True)
    rnd = random.Random(today().toordinal() + 7)

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "انبار"
    _sheet(ws, STOCK_COLS, [[p["code"], f"{p['name']} {p['color']}", p["material"], p["color"], p["width"], p["rolls"],
                             p["roll_meters"], p["meters"], p["opening"], p["buy"], p["sell"], p["supplier"], jstr(p["received"])]
                            for p in data["products"]])
    _sheet(wb.create_sheet("ورود کالا"), STOCK_IN_COLS, [[jstr(r["date"]), r["p"]["code"], f"{r['p']['name']} {r['p']['color']}", r["rolls"],
                                                       r["meters"], r["p"]["supplier"]] for r in data["stock_in"]])
    wb.save(path("stock"))

    rows = []
    bounced = PLANTED["bounced"]
    for s in data["sales"]:
        if s.get("hidden"):
            continue
        name = s["customer"]
        if name == bounced:   # the same customer, typed three ways over the months
            name = rnd.choices([name, "حاج " + name, "حاج " + name.split()[-1]], [.6, .25, .15])[0]
        meters = s["meters"]
        if rnd.random() < .04:   # someone typed a few quantities with Persian digits
            meters = _fa(meters)
        rows.append([s["no"], jstr(s["date"]), name, s["p"]["code"], f"{s['p']['name']} {s['p']['color']}", s["rolls"],
                     meters, s["price"], s["amount"], s["payment"], s["seller"]])
    rows.append(["جمع کل", "", "", "", "", "", "", "", sum(s["amount"] for s in data["sales"] if not s.get("hidden")), "", ""])   # a hand-typed total
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "فروش"
    _sheet(ws, SALES_COLS, rows)
    _sheet(wb.create_sheet(WEB_SHEET), WEB_COLS, [])     # filled by the website's assistant
    wb.save(path("sales"))

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "مشتریان"
    _sheet(ws, CUSTOMER_COLS, [[c["id"], c["name"], c["shop"], c["city"], c["phone"], c["limit"], c["opening"]] for c in data["customers"]])
    _sheet(wb.create_sheet("دریافتی‌ها"), RECEIPT_COLS, [[jstr(r["date"]), r["customer"], r["amount"], r["method"]] for r in data["receipts"]])
    _sheet(wb.create_sheet("چک‌های دریافتی"), CHEQUE_IN_COLS,
           [[q["number"], q["bank"], q["customer"], q["amount"], jstr(q["due"]), q["status"]] for q in sorted(data["cheques_in"], key=lambda q: q["due"])])
    _sheet(wb.create_sheet("چک‌های پرداختی"), CHEQUE_OUT_COLS,
           [[q["number"], q["payee"], q["amount"], jstr(q["due"]), q["status"]] for q in sorted(data["cheques_out"], key=lambda q: q["due"])])
    _sheet(wb.create_sheet("صندوق"), CASH_COLS, [[jstr(r["date"]), r["amount"]] for r in data["deposits"]])
    _sheet(wb.create_sheet("بانک"), BANK_COLS, [[jstr(today()), data["bank"]]])
    wb.save(path("accounts"))

# ---------- reading the workbooks ----------

FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
TITLES = ("حاج آقا", "حاجی", "حاج", "آقای", "آقا", "خانم")

def normal(text):
    """One spelling for comparing names: Persian ی and ک, no titles, single spaces."""
    t = str(text or "").replace("ي", "ی").replace("ك", "ک")
    t = " ".join(t.split())
    for title in TITLES:
        if t.startswith(title + " "):
            t = t[len(title) + 1:]
    return t

def number(value, fixes=None):
    """A cell as a number. Persian digits and separators are accepted; each such fix is counted."""
    if isinstance(value, (int, float)):
        return value
    text = str(value or "").strip()
    plain = text.translate(FA_DIGITS).replace(",", "").replace("٬", "")
    if plain != text and fixes is not None:
        fixes["digits"] = fixes.get("digits", 0) + 1
    try:
        return float(plain) if "." in plain else int(plain)
    except ValueError:
        return None

def jdate(value):
    """'1405/08/03' (any digits) to a date."""
    m = re.match(r"(\d{4})/(\d{1,2})/(\d{1,2})", str(value or "").translate(FA_DIGITS))
    return to_gregorian(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None

def _rows(ws, cols):
    """The sheet's rows as dicts, found by header name, so a moved column still reads right."""
    header = [str(c.value or "").strip() for c in ws[1]]
    missing = [c for c in cols if c not in header]
    if missing:
        raise ValueError(f"برگه «{ws.title}» ستون «{missing[0]}» را ندارد")
    idx = {c: header.index(c) for c in cols}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in row):
            yield {c: row[i] if i < len(row) else None for c, i in idx.items()}

def _book(key):
    try:
        return openpyxl.load_workbook(path(key), data_only=True)
    except Exception:
        raise ValueError(f"فایل {FILES[key][1]} باز نشد. یک فایل اکسل (xlsx) بدهید.")

def _sheet_of(wb, name):
    if name not in wb.sheetnames:
        raise ValueError(f"برگه «{name}» در فایل پیدا نشد")
    return wb[name]

def import_all():
    """Read all three workbooks into the database (the data tables are replaced). Raises ValueError with a
    message for the user when a file can't be read; then nothing is changed."""
    notes = []
    fixes = {}
    stock = _book("stock"); sales_wb = _book("sales"); acc = _book("accounts")

    products = []
    for r in _rows(_sheet_of(stock, "انبار"), STOCK_COLS):
        if not r["کد کالا"]:
            continue
        rolls = number(r["تعداد توپ"], fixes) or 0
        roll_m = number(r["متراژ هر توپ"], fixes) or 0
        meters = number(r["متراژ موجود"], fixes)
        opening = number(r["موجودی اول دوره (متر)"], fixes)
        products.append((str(r["کد کالا"]).strip(), normal(r["شرح کالا"]).replace(" " + normal(r["رنگ"]), "") if r["رنگ"] else normal(r["شرح کالا"]),
                         normal(r["جنس"]), normal(r["رنگ"]), number(r["عرض (سانتی‌متر)"], fixes) or 0, rolls, roll_m,
                         meters if meters is not None else rolls * roll_m, number(r["قیمت خرید هر متر"], fixes) or 0,
                         number(r["قیمت فروش هر متر"], fixes) or 0, normal(r["تأمین‌کننده"]), (jdate(r["تاریخ آخرین ورود"]) or today()).isoformat(),
                         opening or 0))
    codes = {p[0] for p in products}
    stock_in = []
    for r in _rows(_sheet_of(stock, "ورود کالا"), STOCK_IN_COLS):
        d = jdate(r["تاریخ"])
        if d and r["کد کالا"]:
            stock_in.append((d.isoformat(), str(r["کد کالا"]).strip(), number(r["توپ"], fixes) or 0, number(r["متراژ"], fixes) or 0, normal(r["تأمین‌کننده"])))

    customers = []
    for r in _rows(_sheet_of(acc, "مشتریان"), CUSTOMER_COLS):
        cid = number(r["کد مشتری"], fixes)
        if cid is None or not r["نام"]:
            continue
        customers.append((int(cid), normal(r["نام"]), normal(r["فروشگاه"]), normal(r["شهر"]), str(r["تلفن"] or "").translate(FA_DIGITS),
                          number(r["سقف اعتبار"], fixes) or 0, number(r["مانده اول دوره"], fixes) or 0))
    by_name = {c[1]: c[0] for c in customers}
    variants = {}

    def customer_id(raw):
        """The customer a name in another sheet means: the same name, or the one customer whose name contains it."""
        name = normal(raw)
        if name in by_name:
            if name != str(raw).strip():
                variants.setdefault(name, set()).add(str(raw).strip())
            return by_name[name]
        hits = [n for n in by_name if name and (name in n.split() or name in n)]
        if len(hits) == 1:
            variants.setdefault(hits[0], set()).add(str(raw).strip())
            return by_name[hits[0]]
        return None

    invoices, lines, skipped, unknown = {}, [], [], set()
    for r in _rows(_sheet_of(sales_wb, "فروش"), SALES_COLS):
        no = number(r["شماره فاکتور"])
        if not isinstance(no, int):
            skipped.append(str(r["شماره فاکتور"] or "").strip() or "بی‌نام")
            continue
        cid = customer_id(r["مشتری"])
        d = jdate(r["تاریخ"])
        if cid is None or d is None:
            unknown.add(str(r["مشتری"]))
            continue
        meters = number(r["متراژ"], fixes) or 0
        price = number(r["قیمت هر متر"], fixes) or 0
        amount = number(r["مبلغ"], fixes) or round(meters * price)
        inv = invoices.setdefault(no, {"date": d.isoformat(), "customer_id": cid, "seller": normal(r["فروشنده"]),
                                       "payment": normal(r["نوع پرداخت"]) or "نقد", "total": 0})
        inv["total"] += amount
        lines.append((no, str(r["کد کالا"] or "").strip(), number(r["توپ"], fixes) or 0, meters, price, amount))

    receipts = []
    for r in _rows(_sheet_of(acc, "دریافتی‌ها"), RECEIPT_COLS):
        cid, d = customer_id(r["مشتری"]), jdate(r["تاریخ"])
        if cid and d:
            receipts.append((d.isoformat(), cid, number(r["مبلغ"], fixes) or 0, normal(r["روش"])))
    cheques_in = []
    for r in _rows(_sheet_of(acc, "چک‌های دریافتی"), CHEQUE_IN_COLS):
        cid, d = customer_id(r["مشتری"]), jdate(r["تاریخ سررسید"])
        if cid and d:
            cheques_in.append((str(r["شماره چک"] or ""), normal(r["بانک"]), cid, number(r["مبلغ"], fixes) or 0, d.isoformat(), normal(r["وضعیت"])))
    cheques_out = []
    for r in _rows(_sheet_of(acc, "چک‌های پرداختی"), CHEQUE_OUT_COLS):
        d = jdate(r["تاریخ سررسید"])
        if d:
            cheques_out.append((str(r["شماره چک"] or ""), normal(r["در وجه"]), number(r["مبلغ"], fixes) or 0, d.isoformat(), normal(r["وضعیت"])))
    deposits = []
    for r in _rows(_sheet_of(acc, "صندوق"), CASH_COLS):
        d = jdate(r["تاریخ"])
        if d:
            deposits.append((d.isoformat(), number(r["واریز فروش نقدی"], fixes) or 0))
    bank_rows = list(_rows(_sheet_of(acc, "بانک"), BANK_COLS))
    bank = number(bank_rows[-1]["موجودی حساب"], fixes) if bank_rows else 0

    # what was cleaned up, in plain words for the panel
    for name, seen in variants.items():
        others = sorted(seen - {name})
        if others:
            notes.append(("sales", f"نام «{name}» به {len(others)} شکل دیگر هم نوشته شده بود ({'، '.join('«' + o + '»' for o in others)}). همه یک مشتری حساب شد."))
    if fixes.get("digits"):
        notes.append(("sales", f"{fixes['digits']} عدد با رقم فارسی یا جداکننده نوشته شده بود و درست خوانده شد."))
    for s in skipped:
        notes.append(("sales", f"ردیف «{s}» شماره فاکتور نداشت (احتمالاً جمع دستی) و در حساب‌ها نیامد."))
    for u in sorted(unknown):
        notes.append(("sales", f"مشتری «{u}» در فهرست مشتریان پیدا نشد و ردیف‌هایش کنار گذاشته شد."))
    missing = sorted({l[1] for l in lines} - codes)
    if missing:
        notes.append(("stock", f"{len(missing)} کد کالا در فروش هست که در انبار نیست: {'، '.join(missing[:5])}"))

    with store.conn() as c:
        for t in store.DATA_TABLES:
            c.execute(f"DELETE FROM {t}")
        c.executemany("""INSERT INTO products (code, name, material, color, width, rolls, roll_meters, meters, buy_price, sell_price,
                         supplier, received, opening) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""", products)
        c.executemany("INSERT INTO stock_in (date, code, rolls, meters, supplier) VALUES (?,?,?,?,?)", stock_in)
        c.executemany("INSERT INTO cash_deposits (date, amount) VALUES (?,?)", deposits)
        c.executemany("INSERT INTO customers VALUES (?,?,?,?,?,?,?)", customers)
        c.executemany("INSERT INTO invoices (id, date, customer_id, seller, payment, total, source) VALUES (?,?,?,?,?,?, 'excel')",
                      [(no, i["date"], i["customer_id"], i["seller"], i["payment"], i["total"]) for no, i in invoices.items()])
        c.executemany("INSERT INTO invoice_lines (invoice_id, product_code, rolls, meters, price, amount) VALUES (?,?,?,?,?,?)", lines)
        c.executemany("INSERT INTO receipts (date, customer_id, amount, method) VALUES (?,?,?,?)", receipts)
        c.executemany("INSERT INTO cheques_in (number, bank, customer_id, amount, due, status) VALUES (?,?,?,?,?,?)", cheques_in)
        c.executemany("INSERT INTO cheques_out (number, payee, amount, due, status) VALUES (?,?,?,?,?)", cheques_out)
        c.executemany("INSERT INTO import_notes (file, note) VALUES (?,?)", notes)
        store.set_meta(c, "bank", bank)
        store.set_meta(c, "imported_at", store.stamp())
    return {"products": len(products), "customers": len(customers), "invoices": len(invoices), "notes": [n for _, n in notes]}

# ---------- building the demo ----------

def build():
    """Fresh demo: users for the three roles, three new workbooks, and the database read from them."""
    store.init()
    s = business.current()
    with store.conn() as c:
        for t in ("users", "sessions", "reminders", "events", "web_orders", "web_chats"):
            c.execute(f"DELETE FROM {t}")
        for username, role in (("sahebkar", "owner"), ("forooshande", "seller"), ("hesabdar", "accountant")):
            c.execute("INSERT INTO users (username, password, full_name, role) VALUES (?,?,?,?)",
                      (username, store.school_db.hash_password("1234"), business.person(role, s), role))
        store.set_meta(c, "built_for", today().isoformat())
    write_workbooks(simulate())
    result = import_all()
    with store.conn() as c:   # the log starts with the files being read, and this morning's invoices from the sales file
        store.log(c, None, "import", "سه فایل اکسل خوانده شد", {"files": ["انبار", "فروش", "حسابداری"]})
        seller = {"full_name": business.person("seller", s), "role": "seller"}
        for inv in c.execute("SELECT i.id, i.total, cu.name FROM invoices i JOIN customers cu ON cu.id = i.customer_id WHERE i.date = ? ORDER BY i.id",
                             (today().isoformat(),)).fetchall():
            store.log(c, seller, "invoice", f"فاکتور {inv['id']} برای {inv['name']} در فایل فروش ثبت شده", {"files": ["فروش", "انبار"]})
    return result

# ---------- a sale from the phone desk goes into the files ----------

def add_invoice(no, d, customer, lines, payment, seller, cheque=None):
    """lines: [{code, title, rolls, meters, price, amount}]. Appends to sales, takes the rolls out of stock,
    and adds the cheque to accounts. The database is updated by the caller in the same step."""
    wb = openpyxl.load_workbook(path("sales"))
    ws = wb["فروش"]
    total_row = None
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value or "").strip() == "جمع کل":
            total_row = row[0].row
    new = [[no, jstr(d), customer, l["code"], l["title"], l["rolls"], l["meters"], l["price"], l["amount"], payment, seller] for l in lines]
    if total_row:   # keep the hand-typed total at the bottom, as a person would
        ws.insert_rows(total_row, len(new))
        for i, r in enumerate(new):
            for j, v in enumerate(r, 1):
                ws.cell(total_row + i, j, v)
        money = SALES_COLS.index("مبلغ") + 1
        cell = ws.cell(total_row + len(new), money)
        if isinstance(cell.value, (int, float)):
            cell.value += sum(l["amount"] for l in lines)
    else:
        for r in new:
            ws.append(r)
    for col in ("قیمت هر متر", "مبلغ"):
        letter = get_column_letter(SALES_COLS.index(col) + 1)
        for cell in ws[letter][1:]:
            cell.number_format = "#,##0"
    wb.save(path("sales"))

    wb = openpyxl.load_workbook(path("stock"))
    ws = wb["انبار"]
    head = [c.value for c in ws[1]]
    ci, ri, mi = head.index("کد کالا"), head.index("تعداد توپ"), head.index("متراژ موجود")
    for row in ws.iter_rows(min_row=2):
        for l in lines:
            if str(row[ci].value) == l["code"]:
                row[mi].value = max(0, (number(row[mi].value) or 0) - l["meters"])
                row[ri].value = max(0, (number(row[ri].value) or 0) - l["rolls"])
    wb.save(path("stock"))

    if cheque:
        wb = openpyxl.load_workbook(path("accounts"))
        wb["چک‌های دریافتی"].append([cheque["number"], cheque["bank"], customer, cheque["amount"], jstr(cheque["due"]), "در جریان"])
        wb.save(path("accounts"))

LATEST_LAST = ("فروش", "دریافتی‌ها", "چک‌های دریافتی")

def preview(key, limit=60):
    """The workbook's sheets for the panel: header and the first rows of each, plus the row count."""
    wb = _book(key)
    out = []
    for ws in wb.worksheets:
        rows = [[("" if v is None else v) for v in r] for r in ws.iter_rows(values_only=True)]
        latest = ws.title in LATEST_LAST and len(rows) - 1 > limit   # long running lists: the newest rows, where new ones land
        body = rows[-limit:] if latest and limit else rows[1:limit + 1]
        out.append({"name": ws.title, "header": rows[0] if rows else [], "rows": body, "total": max(len(rows) - 1, 0), "latest": latest})
    return out

# ---------- orders from the website ----------

def add_web_order(no, d, time, name, phone, lines, status):
    """An order the website's assistant took goes into the sales file, on its own sheet, as not paid."""
    wb = openpyxl.load_workbook(path("sales"))
    ws = wb[WEB_SHEET] if WEB_SHEET in wb.sheetnames else None
    if ws is None:
        ws = wb.create_sheet(WEB_SHEET)
        _sheet(ws, WEB_COLS, [])
    for l in lines:
        ws.append([no, jstr(d), time, name, phone, l["code"], l["title"], l["rolls"], l["meters"], l["price"], l["amount"], status])
    for col in ("قیمت هر متر", "مبلغ"):
        letter = get_column_letter(WEB_COLS.index(col) + 1)
        for cell in ws[letter][1:]:
            cell.number_format = "#,##0"
    wb.save(path("sales"))

def set_web_order_status(no, status):
    wb = openpyxl.load_workbook(path("sales"))
    if WEB_SHEET in wb.sheetnames:
        ws = wb[WEB_SHEET]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == no:
                row[len(WEB_COLS) - 1].value = status
        wb.save(path("sales"))
